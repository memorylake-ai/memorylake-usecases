#!/usr/bin/env python3
"""Regenerate the architecture-review deck (.pptx) in repos/ledger-service/docs/reviews/.

The committed .pptx is what the demo uploads; you only need this if you edit the slides below.
Unlike demo.py it needs one third-party package:  pip install python-pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
TARGET = HERE / "repos" / "ledger-service" / "docs" / "reviews" / "2025-11-ledger-architecture-review.pptx"

SLIDES = [
    ("Ledger service — architecture review",
     ["Tidewell payments platform · review held 2025-11-18",
      "Presented by Priya Raman · reviewers: platform, SRE, checkout"]),
    ("Event store",
     ["Postgres is the source of truth: append-only ledger_entries table, partitioned by month",
      "Amounts as integer minor units with an ISO 4217 currency column",
      "Idempotency keys for POST requests kept for 72 hours"]),
    ("Rejected: Kafka as the source of truth",
     ["Decision on 2025-11-18: Kafka will not be the ledger's source of truth",
      "Reason 1: the on-call rotation cannot operate a Kafka cluster",
      "Reason 2: exactly-once delivery into the ledger needs transactions we do not want to own"]),
    ("Chosen: transactional outbox",
     ["Events written to ledger_outbox in the same transaction as the entry",
      "Relay worker publishes every 2 seconds; consumers must be idempotent",
      "Recorded as ADR-0004"]),
]


def main() -> None:
    deck = Presentation()
    for title, bullets in SLIDES:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "\n".join(bullets)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    deck.save(TARGET)
    print(f"wrote {TARGET.relative_to(HERE)}")


if __name__ == "__main__":
    main()
