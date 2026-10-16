#!/usr/bin/env python3
"""Regenerate the review decks in data/ from data/campaign.json (needs python-pptx; the demo itself does not).

For every chart it writes what the team's slide tool saves:

  deck-<x>.pptx        one slide: the title and a native column chart, data labels on every bar (the chart's data
                       is stored in the file: in the chart XML's cache and in the embedded workbook)
  deck-<x>-notes.pptx  the same slide, with every number also typed into the speaker notes
  deck-<x>-slide.png   a picture of the slide for the web page (never uploaded; the demo stores the .pptx files)

Larkhaven Outdoor and its numbers are fictional.

    pip install python-pptx pillow && python3 data/make_decks.py
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
FONTS = ["/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
COLOURS = [(68, 114, 196), (237, 125, 49)]  # the slide tool's default series colours


def font(size: int):
    for p in FONTS:
        try:
            return ImageFont.truetype(p, size)
        except OSError:
            pass
    return ImageFont.load_default()


def preview(f: dict, path: Path) -> None:
    """A picture of the slide, drawn from the same data (only for the web page)."""
    W, H = 1280, 720
    im = Image.new("RGB", (W, H), "white")
    g = ImageDraw.Draw(im)
    fs, ft = font(17), font(30)
    g.text((60, 34), f["title"], fill="black", font=ft)
    top = max(v for vs in f["values"].values() for v in vs)
    ymax = 100 if top > 50 else 6
    step = 20 if ymax == 100 else 1
    x0, y0, x1, y1 = 110, 150, 1220, 620
    for t in range(0, ymax + 1, step):
        y = y1 - (y1 - y0) * t / ymax
        g.line((x0, y, x1, y), fill=(225, 225, 225))
        g.text((50, y - 10), (f"${t}" if f["unit"] == "$" else f"{t}%"), fill=(90, 90, 90), font=fs)
    g.line((x0, y1, x1, y1), fill=(150, 150, 150))
    cats, n = list(f["values"]), len(f["series"])
    gw = (x1 - x0) / len(cats)
    bw = (gw - 40) / n
    for i, c in enumerate(cats):
        for k, v in enumerate(f["values"][c]):
            bx = x0 + i * gw + 20 + k * bw
            by = y1 - (y1 - y0) * v / ymax
            g.rectangle((bx, by, bx + bw - 6, y1), fill=COLOURS[k])
            g.text((bx + 2, by - 24), f"${v:.2f}" if f["unit"] == "$" else f"{v:.1f}", fill="black", font=fs)
        tw = g.textlength(c, font=fs)
        g.text((x0 + i * gw + (gw - tw) / 2, y1 + 12), c, fill="black", font=fs)
    if n > 1:
        for k, name in enumerate(f["series"]):
            g.rectangle((520 + k * 200, 100, 536 + k * 200, 116), fill=COLOURS[k])
            g.text((544 + k * 200, 97), name, fill="black", font=fs)
    im.save(path)


def notes_text(f: dict) -> str:
    unit = f["unit"]
    def fmt(v: float) -> str:
        return f"${v:.2f}" if unit == "$" else f"{v:.1f}%"
    parts = [f"{c} " + " / ".join(fmt(v) for v in vs) for c, vs in f["values"].items()]
    order = " / ".join(f["series"])
    return (f"Chart data, for anyone searching later ({order}): " + "; ".join(parts) + ".")


def deck(d: dict, f: dict, notes: bool) -> Path:
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    # The default template carries its own author and dates; the deck's properties are the team's.
    cp = p.core_properties
    cp.author = cp.last_modified_by = f"{d['company']} {d['team']}"
    cp.title, cp.comments, cp.revision = f["title"], "", 1
    cp.created = cp.modified = datetime(2026, 10, 2, 9, 0, 0)
    s = p.slides.add_slide(p.slide_layouts[5])  # title only
    s.shapes.title.text = f["title"]
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
    cd = CategoryChartData()
    cd.categories = list(f["values"])
    for k, name in enumerate(f["series"]):
        cd.add_series(name, [vs[k] for vs in f["values"].values()])
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.7),
                               cd).chart
    chart.has_legend = len(f["series"]) > 1
    if chart.has_legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = '"$"0.00' if f["unit"] == "$" else '0.0'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(12)
    if notes:
        s.notes_slide.notes_text_frame.text = notes_text(f)
    out = HERE / (f["file"] + ("-notes" if notes else "") + ".pptx")
    p.save(out)
    return out


def main() -> None:
    d = json.loads((HERE / "campaign.json").read_text(encoding="utf-8"))
    for f in d["figures"]:
        for notes in (False, True):
            print("wrote", deck(d, f, notes).name)
        preview(f, HERE / (f["file"] + "-slide.png"))
        print("wrote", f["file"] + "-slide.png")


if __name__ == "__main__":
    main()
