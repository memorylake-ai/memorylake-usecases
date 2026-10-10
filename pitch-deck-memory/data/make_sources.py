#!/usr/bin/env python3
"""Regenerate the decks (.pptx) and debriefs (.docx) in sources/ from decks.json.

The committed files are what the demo uploads; you only need this if you edit decks.json.
Unlike demo.py it needs two third-party packages:  pip install python-pptx python-docx
Text is ASCII on purpose: MemoryLake's conflict excerpts mis-decode non-ASCII characters.
"""
import json
from pathlib import Path

import docx
from pptx import Presentation

HERE = Path(__file__).resolve().parent


def main() -> None:
    decks = json.loads((HERE / "decks.json").read_text(encoding="utf-8"))
    out = HERE / "sources"
    out.mkdir(exist_ok=True)
    for name, d in decks.items():
        if d["kind"] == "deck":
            prs = Presentation()
            for title, body in d["slides"]:
                s = prs.slides.add_slide(prs.slide_layouts[1])  # "Title and Content"
                s.shapes.title.text = title
                s.placeholders[1].text = body
            prs.core_properties.title = name
            prs.save(out / name)
        else:
            doc = docx.Document()
            doc.add_heading(d["title"], 0)
            for p in d["paragraphs"]:
                doc.add_paragraph(p)
            doc.core_properties.title = d["title"]
            doc.save(out / name)
        print("wrote", out / name)


if __name__ == "__main__":
    main()
