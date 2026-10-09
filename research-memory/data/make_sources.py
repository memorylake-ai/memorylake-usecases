#!/usr/bin/env python3
"""
Regenerates the PDF and Excel files in data/sources/ from the text below.

You do not need to run this: the generated files are committed. It is here so the
corpus is reviewable as text and easy to change. It needs two packages the demo
itself does not:  pip install reportlab openpyxl
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table

OUT = Path(__file__).resolve().parent / "sources"
STYLES = getSampleStyleSheet()

PAPERS = {
    "halvorsen-2026-lab-study.pdf": dict(
        title="Sodium-ion cells for stationary storage: cost and cycle life in controlled testing",
        byline="Halvorsen Battery Lab, working paper HBL-26-04, May 2026. Funded by Corvane Cells.",
        sections=[
            ("Abstract",
             "We tested 48 prismatic sodium-ion cells (layered-oxide cathode, hard-carbon anode) under a "
             "stationary-storage duty cycle of one full cycle per day at 25 °C. Cells reached 80 percent of "
             "initial capacity after 4,200 cycles on average. At a production volume of 20 GWh a year we "
             "estimate a 2026 pack cost of 61 USD per kWh, 18 percent below LFP packs at 74 USD per kWh."),
            ("1. Method",
             "Cells were cycled at 0.5C charge and discharge between 1.5 V and 4.0 V in a climate chamber held "
             "at 25 °C. Capacity was checked every 100 cycles. The study was funded by the cell maker Corvane "
             "Cells, which supplied all 48 cells."),
            ("2. Results",
             "Mean cycle life to 80 percent capacity was 4,200 cycles (range 3,850 to 4,610). Round-trip "
             "efficiency was 92 percent. Capacity loss was linear after the first 200 cycles."),
            ("3. Cost",
             "Bill of materials at 20 GWh a year gives a 2026 pack cost of 61 USD per kWh against 74 USD per kWh "
             "for LFP. The advantage comes from sodium carbonate and aluminium current collectors on both sides."),
            ("4. Limitations",
             "All tests ran at a constant 25 °C. We did not test partial-state-of-charge cycling or cold climates."),
        ],
    ),
    "gridfield-2026-field-report.pdf": dict(
        title="Sodium-ion in the field: eighteen months of grid storage data",
        byline="Gridfield Analytics, field report GF-2026-09, August 2026. Independent; no vendor funding.",
        sections=[
            ("Summary",
             "We collected operating data from 312 grid storage sites in three climate zones that installed "
             "sodium-ion systems between 2024 and 2025. Projected cycle life to 80 percent capacity is 2,900 "
             "cycles, well below the 4,000-plus cycles reported by laboratory studies."),
            ("1. Why field numbers are lower",
             "Field sites cycle at partial state of charge, see temperatures from -15 °C to 41 °C and run "
             "1.4 cycles a day on average. Sites in the cold zone degraded fastest: 2,300 cycles projected."),
            ("2. Availability",
             "Fleet availability was 97.1 percent. Most outages were inverter faults, not cell failures."),
            ("3. Implications for models",
             "Analysts who use laboratory cycle-life figures will overstate lifetime energy throughput by about "
             "30 percent and understate levelised cost of storage accordingly."),
        ],
    ),
    "storage-policy-brief-q3.pdf": dict(
        title="Policy brief: storage incentives after 2027",
        byline="Larkspur Research policy desk, Q3 2026.",
        sections=[
            ("Summary",
             "From 1 January 2027 the storage investment credit adds a 10 percent bonus for systems whose cells "
             "are made domestically. No sodium-ion cell plant is expected to be producing domestically before "
             "2028, so the bonus favours LFP for at least the first year."),
            ("Timeline",
             "Rule-making closes in November 2026. The first credits under the new rule are claimable for "
             "projects that start construction after 1 January 2027."),
        ],
    ),
}

COSTS = [
    ("Chemistry", "Year", "Pack cost (USD/kWh)", "Source"),
    ("LFP", 2025, 79, "Larkspur tracker"),
    ("LFP", 2026, 74, "Larkspur tracker"),
    ("LFP", 2027, 70, "Larkspur forecast"),
    ("Sodium-ion", 2025, 72, "Larkspur tracker"),
    ("Sodium-ion", 2026, 61, "Halvorsen HBL-26-04"),
    ("Sodium-ion", 2027, 55, "Larkspur forecast"),
]
ASSUMPTIONS = [
    ("Parameter", "Value", "Set on", "Note"),
    ("Sodium-ion cycle life (cycles to 80%)", 2900, "2026-09-15", "Was 4,200 (lab) until the Q3 review"),
    ("LFP cycle life (cycles to 80%)", 6000, "2026-07-14", "Field data agrees with lab"),
    ("Cycles per day", 1.2, "2026-07-14", ""),
    ("Discount rate", 0.08, "2026-07-14", ""),
    ("Vendor-funded study haircut", 0.20, "2026-07-14", "Mira's rule: discount vendor-funded results by 20%"),
]


def paper(name: str, spec: dict) -> None:
    doc = SimpleDocTemplate(str(OUT / name), pagesize=A4, title=spec["title"], author=spec["byline"])
    flow = [Paragraph(spec["title"], STYLES["Title"]), Paragraph(spec["byline"], STYLES["Italic"]), Spacer(1, 12)]
    for i, (heading, body) in enumerate(spec["sections"]):
        if i == 3:
            flow.append(PageBreak())
        flow += [Paragraph(heading, STYLES["Heading2"]), Paragraph(body, STYLES["BodyText"]), Spacer(1, 8)]
    if name.startswith("halvorsen"):
        flow += [Paragraph("Table 1. Cycle life by cell batch", STYLES["Heading3"]),
                 Table([["Batch", "Cells", "Cycles to 80%"], ["A", 16, 4180], ["B", 16, 4610], ["C", 16, 3850]])]
    doc.build(flow)


def workbook() -> None:
    wb = Workbook()
    for title, rows in (("Pack costs", COSTS), ("Assumptions", ASSUMPTIONS)):
        ws = wb.active if title == "Pack costs" else wb.create_sheet()
        ws.title = title
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in "ABCD":
            ws.column_dimensions[col].width = 38 if col in "AD" else 16
    wb.save(OUT / "storage-cost-model.xlsx")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    for name, spec in PAPERS.items():
        paper(name, spec)
    workbook()
    print("\n".join(sorted(p.name for p in OUT.iterdir())))
